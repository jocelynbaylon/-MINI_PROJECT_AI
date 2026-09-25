"""
export_engine.py
-----------------
Queries SQLite by course_code (via db_manager) and compiles the record into
the official CCS syllabus layout using Jinja2, producing a browser-ready
HTML file. export_docx() converts the same data into a formatted Word
document using python-docx.
"""

import os
import base64
from jinja2 import Environment, FileSystemLoader

import db_manager

BASE_DIR = os.path.dirname(__file__)
TEMPLATE_DIR = os.path.join(BASE_DIR, "templates")
LOGO_PATH = os.path.join(BASE_DIR, "ccs.png")

# Built once at import time and reused for every render: recreating the Jinja
# Environment (which re-reads and re-parses the template file) on every single
# export/preview call was pure overhead, since the template never changes at
# runtime.
_ENV = Environment(loader=FileSystemLoader(TEMPLATE_DIR))
_TEMPLATE = _ENV.get_template("uphsd_ccs_template.html")

# The logo PNG is likewise constant, so read + base64-encode it once instead
# of hitting disk again on every export (including the auto-preview that
# fires after every "Generate Syllabus" click).
_logo_cache = {"value": None}


def _logo_b64() -> str:
    if _logo_cache["value"] is None:
        if os.path.exists(LOGO_PATH):
            with open(LOGO_PATH, "rb") as f:
                _logo_cache["value"] = base64.b64encode(f.read()).decode()
        else:
            _logo_cache["value"] = ""
    return _logo_cache["value"]


def render_html(course_code: str) -> str:
    """Returns the rendered syllabus as an HTML string."""
    data = db_manager.get_syllabus(course_code)
    return _TEMPLATE.render(
        course=data["course"],
        course_outcomes=data["course_outcomes"],
        weekly_schedule=data["weekly_schedule"],
        logo_b64=_logo_b64(),
    )


def export_html(course_code: str, out_path: str) -> str:
    html = render_html(course_code)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html)
    return out_path


def export_docx(course_code: str, out_path: str) -> str:
    """Converts the syllabus data into a beautiful DOCX file using python-docx."""
    import docx
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml.ns import nsdecls
    from docx.oxml import parse_xml

    data = db_manager.get_syllabus(course_code)
    course = data["course"]
    
    document = docx.Document()
    
    # Set default font
    style = document.styles['Normal']
    style.font.name = 'Segoe UI'
    style.font.size = Pt(10)
    
    def set_col_widths(table, widths):
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
                
    def shade_cells(cells, color="F6F3EE"):
        for cell in cells:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{color}"/>')
            tcPr.append(shd)
            
    def remove_borders(table):
        for row in table.rows:
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
                tcPr.append(tcBorders)

    # Header using borderless table for alignment
    hdr_table = document.add_table(rows=1, cols=2)
    hdr_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_borders(hdr_table)
    set_col_widths(hdr_table, [3.0, 3.5])
    
    cell_logo = hdr_table.cell(0, 0)
    p_logo = cell_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if os.path.exists(LOGO_PATH):
        r_logo = p_logo.add_run()
        r_logo.add_picture(LOGO_PATH, width=Inches(0.6))
        
    cell_title = hdr_table.cell(0, 1)
    cell_title.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p_title = cell_title.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = p_title.add_run("College of Computer Studies")
    run_title.font.name = 'Segoe UI'
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x7A, 0x0C, 0x1E)
    
    p_sub = document.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run(f"{course['course_code']}: {course['course_title']}")
    run_sub.font.name = 'Segoe UI'
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    
    p_meta = document.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = p_meta.add_run(f"{course['semester']} · SY {course['school_year']} · Section: {course['section']} · Instructor: {course['instructor']}")
    run_meta.font.name = 'Segoe UI'
    run_meta.font.size = Pt(9)
    run_meta.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    
    document.add_paragraph()
    
    def add_heading(text):
        p = document.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x7A, 0x0C, 0x1E)
        
    # 1. Course Learning Outcomes
    add_heading("1. Course Learning Outcomes (CLOs)")
    table = document.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_texts = ["CLO #", "Bloom's Level", "Description", "Mapped PLOs"]
    shade_cells(hdr_cells, "F6F3EE")
    for i, t in enumerate(hdr_texts):
        hdr_cells[i].text = t
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0x7A, 0x0C, 0x1E)
    set_col_widths(table, [0.8, 1.2, 3.5, 1.0])
        
    for co in data["course_outcomes"]:
        row_cells = table.add_row().cells
        row_cells[0].text = f"CLO {co['clo_number']}"
        row_cells[0].paragraphs[0].runs[0].font.bold = True
        row_cells[1].text = co['bloom_level']
        row_cells[2].text = co['co_description']
        plos = ", ".join([f"PLO {p}" for p in str(co['mapped_po']).split(',')])
        row_cells[3].text = plos
        
    # 2. Grading Breakdown
    add_heading("2. Grading System & Breakdown")
    table2 = document.add_table(rows=1, cols=2)
    table2.style = 'Table Grid'
    hdr_cells2 = table2.rows[0].cells
    shade_cells(hdr_cells2, "F6F3EE")
    hdr_cells2[0].text = "Component"
    hdr_cells2[1].text = "Weight / Share"
    for c in hdr_cells2:
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x7A, 0x0C, 0x1E)
    set_col_widths(table2, [2.5, 4.0])
        
    row_cells = table2.add_row().cells
    row_cells[0].text = f"Class Standing ({course.get('class_standing_weight', 0)}%)"
    row_cells[0].paragraphs[0].runs[0].font.bold = True
    row_cells[1].text = f"Quizzes ({course.get('quizzes_pct', 0)}%) | Research ({course.get('research_pct', 0)}%) | Seatwork/Lab ({course.get('seatwork_lab_pct', 0)}%)"
    
    row_cells = table2.add_row().cells
    row_cells[0].text = "Major Examination"
    row_cells[0].paragraphs[0].runs[0].font.bold = True
    row_cells[1].text = f"{course.get('major_exam_weight', 0)}%"
    
    # 3. Weekly Schedule
    add_heading("3. 18-Week Course Schedule & Learning Plan")
    table3 = document.add_table(rows=1, cols=6)
    table3.style = 'Table Grid'
    hdr_cells3 = table3.rows[0].cells
    shade_cells(hdr_cells3, "F6F3EE")
    headers3 = ["Wk", "Topics", "Lesson Learning Outcomes (LLOs)", "Teaching & Learning Activities", "Assessment & Evidence", "Aligned CLO"]
    for i, t in enumerate(headers3):
        hdr_cells3[i].text = t
        hdr_cells3[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells3[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0x7A, 0x0C, 0x1E)
    set_col_widths(table3, [0.5, 1.2, 2.0, 1.2, 1.0, 0.6])
        
    for wk in data["weekly_schedule"]:
        row_cells = table3.add_row().cells
        row_cells[0].text = f"Wk {wk['week_number']}\n[{wk['period']}]"
        row_cells[0].paragraphs[0].runs[0].font.bold = True
        row_cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        
        row_cells[1].text = ", ".join(wk['topics'])
        row_cells[1].paragraphs[0].runs[0].font.bold = True
        
        llos_text = "\n".join([f"[{l['category']}] {l['outcome_text']}" for l in wk['llos']])
        row_cells[2].text = llos_text
        
        row_cells[3].text = wk['teaching_learning_activity']
        row_cells[4].text = f"{wk['assessment_tool']}\n{wk['evidence']}"
        row_cells[5].text = ", ".join([f"CLO {c}" for c in wk['aligned_co']])
        
    document.save(out_path)
    return out_path


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Usage: python export_engine.py <COURSE_CODE> [html|docx]")
        sys.exit(1)
    code = sys.argv[1]
    fmt = sys.argv[2] if len(sys.argv) > 2 else "html"
    out = f"{code.replace(' ', '_')}_Syllabus.{fmt}"
    if fmt == "docx":
        export_docx(code, out)
    else:
        export_html(code, out)
    print(f"Exported -> {out}")